import io
import pytest
import openpyxl
from fastapi.testclient import TestClient
from backend.main import app
from backend.excel_parser import parse_excel_floorplan, ExcelParsingError
from backend.geometry import generate_excel_template_bytes, process_geometry_solve

client = TestClient(app)


def test_parse_default_excel_template():
    """Test parsing the standardized generated Excel template."""
    template_bytes = generate_excel_template_bytes()
    rooms, openings, hvac = parse_excel_floorplan(template_bytes)

    assert len(rooms) == 2
    assert len(openings) == 3
    
    living_room = next(r for r in rooms if r.name == "Living Room")
    assert len(living_room.points) == 4
    assert living_room.fillet_radius == 0.0
    assert living_room.points[0] == [0.0, 0.0]
    assert living_room.points[1] == [6000.0, 0.0]
    assert living_room.points[2] == [6000.0, 5000.0]
    assert living_room.points[3] == [0.0, 5000.0]

    corridor = next(r for r in rooms if r.name == "L-Shaped Corridor/Balcony")
    assert len(corridor.points) == 4
    assert corridor.fillet_radius == 500.0
    assert corridor.points[0] == [0.0, 5000.0]
    assert corridor.points[1] == [3000.0, 5000.0]
    assert corridor.points[2] == [3000.0, 8000.0]
    assert corridor.points[3] == [0.0, 8000.0]


def test_parse_and_solve_pipeline():
    """Verify end-to-end solving of an uploaded template."""
    template_bytes = generate_excel_template_bytes()
    rooms, openings, hvac = parse_excel_floorplan(template_bytes)
    solution = process_geometry_solve(rooms, openings)

    assert len(solution.rooms) == 2
    assert len(solution.wall_polygons) >= 1
    assert len(solution.dimensions) >= 3
    assert len(solution.doors) == 1
    assert len(solution.windows) == 2
    
    lr_result = next(r for r in solution.rooms if r.name == "Living Room")
    assert lr_result.carpet_area_sqm == 30.0


def test_corrupt_file_raises_error():
    """Verify corrupt non-Excel data raises ExcelParsingError."""
    with pytest.raises(ExcelParsingError) as exc_info:
        parse_excel_floorplan(b"NOT_A_VALID_EXCEL_STREAM")
    assert "Failed to read Excel workbook" in str(exc_info.value) or "empty" in str(exc_info.value)


def test_missing_headers_raises_error():
    """Verify missing required columns raises ExcelParsingError."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Foo", "Bar", "Baz"])
    ws.append([1, 2, 3])
    
    buf = io.BytesIO()
    wb.save(buf)
    
    with pytest.raises(ExcelParsingError) as exc_info:
        parse_excel_floorplan(buf.getvalue())
    assert "Could not find required" in str(exc_info.value)


def test_room_with_fewer_than_three_points():
    """Verify rooms with fewer than 3 vertices fail geometry validation."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Room Name", "Fillet Radius (mm)", "Point Index", "X Coordinate (mm)", "Y Coordinate (mm)"])
    ws.append(["Bad Room", 0, 1, 0.0, 0.0])
    ws.append(["Bad Room", 0, 2, 1000.0, 0.0])
    
    buf = io.BytesIO()
    wb.save(buf)

    with pytest.raises(ExcelParsingError) as exc_info:
        parse_excel_floorplan(buf.getvalue())
    assert "must have at least 3 vertices" in str(exc_info.value)


def test_degenerate_zero_area_polygon():
    """Verify collinear vertices with zero area are rejected."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Room Name", "Fillet Radius (mm)", "Point Index", "X Coordinate (mm)", "Y Coordinate (mm)"])
    ws.append(["Line Room", 0, 1, 0.0, 0.0])
    ws.append(["Line Room", 0, 2, 1000.0, 0.0])
    ws.append(["Line Room", 0, 3, 2000.0, 0.0])
    
    buf = io.BytesIO()
    wb.save(buf)

    with pytest.raises(ExcelParsingError) as exc_info:
        parse_excel_floorplan(buf.getvalue())
    assert "degenerate polygon with zero area" in str(exc_info.value)


def test_api_upload_endpoint_success():
    """Test POST /api/upload-plan endpoint with valid Excel template."""
    template_bytes = generate_excel_template_bytes()
    files = {
        "file": ("floorplan.xlsx", template_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    response = client.post("/api/upload-plan", files=files)
    assert response.status_code == 200
    data = response.json()
    assert "rooms" in data
    assert len(data["rooms"]) == 2
    assert "wall_polygons" in data
    assert "dimensions" in data


def test_api_upload_endpoint_invalid_extension():
    """Test POST /api/upload-plan rejects non-excel files."""
    files = {
        "file": ("notes.txt", b"plain text content", "text/plain")
    }
    response = client.post("/api/upload-plan", files=files)
    assert response.status_code == 400
    assert "Invalid file format" in response.json()["detail"]


def test_api_upload_endpoint_corrupt_file():
    """Test POST /api/upload-plan returns 400 on corrupt excel stream."""
    files = {
        "file": ("broken.xlsx", b"invalid bytes", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    response = client.post("/api/upload-plan", files=files)
    assert response.status_code == 400
    assert "Failed to read Excel workbook" in response.json()["detail"]


def test_api_download_custom_plan_endpoint():
    """Test GET /api/download-custom-plan endpoint returns valid 3-BHK Excel spreadsheet."""
    response = client.get("/api/download-custom-plan")
    assert response.status_code == 200
    assert "Custom_Full_House_Plan.xlsx" in response.headers.get("content-disposition", "")
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.headers.get("content-type", "")
    assert len(response.content) > 1000

    rooms, openings, hvac = parse_excel_floorplan(response.content)
    assert len(rooms) == 6
    assert len(openings) == 12
    assert len(hvac) == 4


def test_parse_full_6_disciplines_roundtrip():
    """Verify bidirectional 6-sheet roundtrip parsing all disciplines."""
    from backend.excel_parser import parse_excel_floorplan_full
    template_bytes = generate_excel_template_bytes()

    # 1. Test full 9-tuple parser
    (
        rooms, openings, hvac,
        furniture, plm_fixtures, plm_pipes,
        el_comps, el_ckts, el_routes
    ) = parse_excel_floorplan_full(template_bytes)

    assert len(rooms) >= 2
    assert len(openings) >= 1
    assert len(hvac) >= 1
    assert len(furniture) >= 2
    assert furniture[0].type in ("Sofa", "Dining Table")
    assert len(plm_fixtures) >= 1
    assert plm_fixtures[0].type == "Wash Basin"
    assert len(plm_pipes) >= 1
    assert len(el_comps) >= 2

    # 2. Test backwards compatibility 3-tuple
    legacy_rooms, legacy_openings, legacy_hvac = parse_excel_floorplan(template_bytes)
    assert len(legacy_rooms) == len(rooms)
    assert len(legacy_openings) == len(openings)
    assert len(legacy_hvac) == len(hvac)

