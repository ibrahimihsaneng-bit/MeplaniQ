import pytest
from fastapi.testclient import TestClient
from backend.main import app
from backend.database import init_db
import openpyxl
import io

client = TestClient(app)

def setup_module(module):
    init_db()

def test_plumbing_e2e_full_flow():
    # 1. Start with an empty project and test Excel export (empty states)
    resp_empty = client.post("/api/export-excel", json={"project_name": "E2E Plumbing Test", "rooms": [], "plumbing_fixtures": [], "plumbing_pipes": []})
    assert resp_empty.status_code == 200
    wb_empty = openpyxl.load_workbook(io.BytesIO(resp_empty.content))
    assert len(wb_empty.sheetnames) >= 5
    assert "Plumbing Schedule" in wb_empty.sheetnames
    
    # Verify the empty state styling on Plumbing Schedule
    ws_plumbing_empty = wb_empty["Plumbing Schedule"]
    assert "Plumbing & Sanitary" in ws_plumbing_empty.cell(row=1, column=1).value
    assert ws_plumbing_empty.cell(row=3, column=1).value == "—"
    assert "No plumbing fixtures" in ws_plumbing_empty.cell(row=3, column=2).value
    
    # 2. Add some rooms and plumbing data
    payload = {
        "project_name": "E2E Plumbing Master",
        "rooms": [
            {"name": "Master Bathroom", "points": [[0, 0], [3000, 0], [3000, 3000], [0, 3000]], "fillet_radius": 0}
        ],
        "plumbing_fixtures": [
            {
                "id": "WC-101",
                "type": "WC",
                "category": "Bathroom fixtures",
                "host_room_id": "Master Bathroom",
                "width": 450,
                "height": 700,
                "x": 1500,
                "y": 500,
                "rotation": 0,
                "connected_to": "SEWER-01",
                "clearance_status": "Clear"
            }
        ],
        "plumbing_pipes": [
            {
                "id": "SEWER-01",
                "type": "Drainage Pipe",
                "host_room_id": "Master Bathroom",
                "diameter": 110,
                "path": [[1500, 500], [1500, -1000]],
                "connected_to": "City Main"
            }
        ]
    }
    
    # 3. Export again and verify data integrity
    resp_full = client.post("/api/export-excel", json=payload)
    assert resp_full.status_code == 200
    assert "E2E_Plumbing_Master_Full_Schedule.xlsx" in resp_full.headers["content-disposition"]
    
    wb_full = openpyxl.load_workbook(io.BytesIO(resp_full.content))
    ws_plumbing_full = wb_full["Plumbing Schedule"]
    
    # Verify Row 3 is WC-101
    assert ws_plumbing_full.cell(row=3, column=1).value == "WC-101"
    assert ws_plumbing_full.cell(row=3, column=2).value == "Bathroom fixtures / WC"
    
    # Verify Row 4 is SEWER-01
    assert ws_plumbing_full.cell(row=4, column=1).value == "SEWER-01"
    assert ws_plumbing_full.cell(row=4, column=2).value == "Water/Drainage lines / Drainage Pipe"
    
    # Verify Run Length is 1500 mm (from [1500, 500] to [1500, -1000])
    assert ws_plumbing_full.cell(row=4, column=7).value == "1500.0 mm"
